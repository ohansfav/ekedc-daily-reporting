import openpyxl
import os
from collections import defaultdict

p = r'C:\Users\ohanu\OneDrive\Desktop\ekedc\data\Eko_Power_App_Cashier_Reporting_Template.xlsx'
print('exists', os.path.exists(p))
if not os.path.exists(p):
    raise SystemExit(1)

wb = openpyxl.load_workbook(p, data_only=True)
print('sheets', wb.sheetnames)

ws = wb['Cashier Roster']
roster = defaultdict(list)
for row in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
    if not row or len(row) < 9:
        continue
    name, district, status = row[1], row[2], row[7]
    if name and district and status == 'Active':
        roster[district].append(name)

print('--- Active roster by district')
for district in sorted(roster):
    print(f'{district}: {roster[district]}')
print('--- Total active cashiers:', sum(len(v) for v in roster.values()))
